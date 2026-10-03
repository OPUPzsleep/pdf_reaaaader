"""Genera los archivos de Office de prueba (tests/fixtures/office/*) con python-docx, openpyxl y python-pptx.
Se guardan en el repositorio: las pruebas no necesitan Python. Uso: python3 tests/fixtures/generar-office.py"""
import datetime, io, os, struct, zlib
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Inches
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as PI, Pt as PPt, Emu

SALIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'office')
os.makedirs(SALIDA, exist_ok=True)

def png(ancho, alto, color_fn):
    """PNG RGB sin dependencias"""
    filas = b''.join(b'\x00' + b''.join(bytes(color_fn(x, y)) for x in range(ancho)) for y in range(alto))
    def trozo(tipo, datos):
        c = struct.pack('>I', len(datos)) + tipo + datos
        return c + struct.pack('>I', zlib.crc32(tipo + datos) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + trozo(b'IHDR', struct.pack('>IIBBBBB', ancho, alto, 8, 2, 0, 0, 0)) + trozo(b'IDAT', zlib.compress(filas)) + trozo(b'IEND', b'')

IMG = png(160, 90, lambda x, y: (int(255 * x / 160), 90, int(255 * y / 90)))
open(os.path.join(SALIDA, 'imagen.png'), 'wb').write(IMG)

def campo(par, instr):
    for tipo in ('begin', None, 'separate', 'text', 'end'):
        r = par.add_run()
        if tipo == 'text':
            r.text = '1'
        elif tipo is None:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' {instr} '; r._r.append(it)
        else:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tipo); r._r.append(fc)

def enlace(par, texto, url):
    rid = par.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    st = OxmlElement('w:rStyle'); st.set(qn('w:val'), 'Hyperlink'); rpr.append(st)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rpr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    r.append(rpr); t = OxmlElement('w:t'); t.text = texto; r.append(t); h.append(r)
    par._p.append(h)

def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color); tcPr.append(shd)

# ───────────── Word ─────────────
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.header.paragraphs[0].text = 'Informe confidencial — Departamento de Ventas'
pie = sec.footer.paragraphs[0]; pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
pie.add_run('Página '); campo(pie, 'PAGE'); pie.add_run(' de '); campo(pie, 'NUMPAGES')
d.add_heading('Informe trimestral de ventas', 0)
d.add_paragraph('Resumen ejecutivo del trimestre con cifras, comparativas y previsiones para el siguiente periodo fiscal.', style='Subtitle')
d.add_heading('1. Introducción', 1)
p = d.add_paragraph('Este documento resume los ')
p.add_run('resultados').bold = True
p.add_run(' del trimestre y las ')
p.add_run('previsiones').italic = True
p.add_run(' para el siguiente periodo. Incluye ')
p.add_run('subrayado').underline = True
p.add_run(', ')
r = p.add_run('color rojo'); r.font.color.rgb = RGBColor(0xC0, 0, 0)
p.add_run(', ')
r = p.add_run('texto grande'); r.font.size = Pt(16)
p.add_run(' y un enlace a '); enlace(p, 'example.com', 'https://example.com'); p.add_run('. Fórmula: H'); s = p.add_run('2'); s.font.subscript = True; p.add_run('O y x'); s = p.add_run('2'); s.font.superscript = True; p.add_run('.')
d.add_paragraph('Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat. Duis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore eu fugiat nulla pariatur.').alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
d.add_heading('2. Listas', 1)
for t in ('Primer punto', 'Segundo punto', 'Tercer punto'): d.add_paragraph(t, style='List Bullet')
d.add_paragraph('Subpunto anidado', style='List Bullet 2')
for t in ('Preparar los datos', 'Revisar las cifras', 'Enviar el informe'): d.add_paragraph(t, style='List Number')
d.add_heading('3. Tabla de resultados', 1)
tabla = d.add_table(rows=4, cols=4); tabla.style = 'Table Grid'
cab = ['Producto', 'Unidades', 'Precio', 'Total']
for i, t in enumerate(cab):
    c = tabla.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(t); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); sombrear(c, '1F3864')
datos = [('Manzanas', '120', '1,50', '180,00'), ('Peras', '80', '2,25', '180,00'), ('Uvas', '45', '3,10', '139,50')]
for fi, fila in enumerate(datos, 1):
    for ci, t in enumerate(fila):
        c = tabla.rows[fi].cells[ci]; c.text = t
        if ci: c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
a = tabla.cell(3, 0).merge(tabla.cell(3, 1)); 
d.add_paragraph()
d.add_heading('4. Imagen', 1)
d.add_picture(io.BytesIO(IMG), width=Cm(8))
d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_heading('5. Índice con tabulador', 1)
for t, n in (('Introducción', '1'), ('Listas', '2'), ('Resultados', '3')):
    p = d.add_paragraph(); p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS); p.add_run(f'{t}\t{n}')
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
d.add_heading('6. Segunda página', 1)
d.add_paragraph('Esta página va después de un salto de página explícito.')
# Sección apaisada
s2 = d.add_section(WD_SECTION.NEW_PAGE); s2.orientation = WD_ORIENT.LANDSCAPE; s2.page_width, s2.page_height = Cm(29.7), Cm(21)
d.add_heading('7. Sección apaisada', 1)
d.add_paragraph('Esta sección usa una página horizontal.')
d.save(os.path.join(SALIDA, 'informe.docx'))

# ───────────── Excel ─────────────
wb = Workbook()
ws = wb.active; ws.title = 'Ventas'
ws.append(['Informe de ventas 2024']); ws.merge_cells('A1:E1'); ws['A1'].font = Font(size=16, bold=True, color='1F3864'); ws['A1'].alignment = Alignment(horizontal='center')
ws.append([])
ws.append(['Producto', 'Fecha', 'Unidades', 'Precio', 'Descuento'])
fino = Side(style='thin', color='444444')
for c in ws[3]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.border = Border(left=fino, right=fino, top=fino, bottom=fino); c.alignment = Alignment(horizontal='center')
filas = [('Manzanas', datetime.date(2024, 1, 15), 120, 1.5, 0.05), ('Peras', datetime.date(2024, 2, 3), 80, 2.25, 0.1), ('Uvas rojas de mesa sin pepitas', datetime.date(2024, 3, 30), 45, 3.1, 0), ('Melones', datetime.date(2024, 4, 12), 1234567, 12.999, 0.255)]
for i, f in enumerate(filas, 4):
    ws.append(list(f))
    ws.cell(i, 2).number_format = 'dd/mm/yyyy'; ws.cell(i, 3).number_format = '#,##0'; ws.cell(i, 4).number_format = '#,##0.00 "€"'; ws.cell(i, 5).number_format = '0.0%'
    for c in ws[i]: c.border = Border(left=fino, right=fino, top=fino, bottom=fino)
    if i % 2: 
        for c in ws[i]: c.fill = PatternFill('solid', fgColor='DDEBF7')
ws.append(['Total', None, '=SUM(C4:C7)']); ws['A8'].font = Font(bold=True)
ws.append(['Notas con texto largo que desborda hacia las celdas vacías de la derecha']); ws['A9'].font = Font(italic=True, color='7F7F7F')
ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 12; ws.column_dimensions['C'].width = 12; ws.column_dimensions['D'].width = 14; ws.column_dimensions['E'].width = 12
ws.row_dimensions[1].height = 28
ws.page_setup.orientation = 'portrait'; ws.sheet_properties.pageSetUpPr = None
ws2 = wb.create_sheet('Resumen')
ws2.append(['Mes', 'Importe', 'Estado'])
for i, (m, v, e) in enumerate((('Enero', 1500.5, True), ('Febrero', -230.25, False), ('Marzo', 0, True)), 2):
    ws2.append([m, v, e]); ws2.cell(i, 2).number_format = '#,##0.00;[Red]-#,##0.00'
ws2['A6'] = 'Texto con ajuste de línea automático en una celda estrecha'; ws2['A6'].alignment = Alignment(wrap_text=True, vertical='top'); ws2.column_dimensions['A'].width = 16; ws2.column_dimensions['B'].width = 12; ws2.column_dimensions['C'].width = 14; ws2.row_dimensions[6].height = 60
ws3 = wb.create_sheet('Ancha')
ws3.append([f'Columna {i}' for i in range(1, 25)])
ws3.append([i * 100.5 for i in range(1, 25)])
for i in range(1, 25): ws3.column_dimensions[chr(64 + i) if i <= 26 else 'A'].width = 11
ws3.page_setup.orientation = 'landscape'
wsg = wb.create_sheet('Gráficos')
wsg.append(['Mes', 'Ventas', 'Gastos'])
for fila in (('Ene', 120, 80), ('Feb', 150, 95), ('Mar', 90, 110), ('Abr', 180, 100), ('May', 210, 130)):
    wsg.append(list(fila))
col = BarChart(); col.type = 'col'; col.title = 'Ventas y gastos'; col.y_axis.title = 'Euros'
col.add_data(Reference(wsg, min_col=2, max_col=3, min_row=1, max_row=6), titles_from_data=True); col.set_categories(Reference(wsg, min_col=1, min_row=2, max_row=6))
col.width, col.height = 14, 7.5
wsg.add_chart(col, 'E2')
lin = LineChart(); lin.title = 'Tendencia'; lin.add_data(Reference(wsg, min_col=2, max_col=3, min_row=1, max_row=6), titles_from_data=True); lin.set_categories(Reference(wsg, min_col=1, min_row=2, max_row=6))
lin.width, lin.height = 14, 7.5
wsg.add_chart(lin, 'E18')
pie = PieChart(); pie.title = 'Reparto'; pie.add_data(Reference(wsg, min_col=2, min_row=1, max_row=6), titles_from_data=True); pie.set_categories(Reference(wsg, min_col=1, min_row=2, max_row=6))
pie.width, pie.height = 6.4, 6
wsg.add_chart(pie, 'A9')
wb.save(os.path.join(SALIDA, 'libro.xlsx'))

# ───────────── PowerPoint ─────────────
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
t = prs.slides.add_slide(prs.slide_layouts[0]); t.shapes.title.text = 'Plan de lanzamiento'; t.placeholders[1].text = 'Producto nuevo · Otoño 2024'
c = prs.slides.add_slide(prs.slide_layouts[1]); c.shapes.title.text = 'Objetivos'
tf = c.placeholders[1].text_frame; tf.text = 'Aumentar las ventas un 20 %'
for txt, nivel in (('Reforzar la marca', 0), ('Campaña en redes sociales', 1), ('Nuevos mercados', 0)):
    pr = tf.add_paragraph(); pr.text = txt; pr.level = nivel
s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = 'Formas y colores'
for i, (forma, col) in enumerate(((MSO_SHAPE.RECTANGLE, 'C00000'), (MSO_SHAPE.OVAL, '2E75B6'), (MSO_SHAPE.ROUNDED_RECTANGLE, '70AD47'), (MSO_SHAPE.RIGHT_ARROW, 'ED7D31'), (MSO_SHAPE.ISOSCELES_TRIANGLE, '7030A0'))):
    sh = s.shapes.add_shape(forma, PI(0.8 + i * 2.4), PI(2.2), PI(2), PI(1.5)); sh.fill.solid(); sh.fill.fore_color.rgb = PRGB.from_string(col); sh.text_frame.text = f'Forma {i + 1}'; sh.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
tb = s.shapes.add_textbox(PI(0.8), PI(4.4), PI(8), PI(1.2)); tf = tb.text_frame; tf.word_wrap = True
tf.text = 'Texto en negrita, '; tf.paragraphs[0].runs[0].font.bold = True; tf.paragraphs[0].runs[0].font.size = PPt(24)
r = tf.paragraphs[0].add_run(); r.text = 'cursiva roja'; r.font.italic = True; r.font.color.rgb = PRGB(0xC0, 0, 0); r.font.size = PPt(24)
p2 = tf.add_paragraph(); p2.text = 'Centrado y pequeño'; p2.alignment = PP_ALIGN.CENTER; p2.runs[0].font.size = PPt(14)
pic = prs.slides.add_slide(prs.slide_layouts[5]); pic.shapes.title.text = 'Imagen y tabla'
pic.shapes.add_picture(io.BytesIO(IMG), PI(0.8), PI(1.8), width=PI(4.5))
tbl = pic.shapes.add_table(3, 3, PI(6), PI(1.8), PI(6.5), PI(2)).table
for r_i, fila in enumerate((('Región', 'Ventas', 'Cuota'), ('Norte', '1.250', '35 %'), ('Sur', '980', '27 %'))):
    for c_i, txt in enumerate(fila): tbl.cell(r_i, c_i).text = txt
fin = prs.slides.add_slide(prs.slide_layouts[6]); tb = fin.shapes.add_textbox(PI(3), PI(3), PI(7), PI(1.5)); tb.text_frame.text = '¡Gracias!'; tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER; tb.text_frame.paragraphs[0].runs[0].font.size = PPt(54)
gs = prs.slides.add_slide(prs.slide_layouts[5]); gs.shapes.title.text = 'Gráficos'
cd = CategoryChartData(); cd.categories = ['Norte', 'Sur', 'Este', 'Oeste']; cd.add_series('2023', (19.2, 21.4, 16.7, 12.9)); cd.add_series('2024', (22.3, 18.2, 21.1, 15.6))
g1 = gs.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PI(0.6), PI(1.6), PI(6), PI(4.6), cd).chart
g1.has_legend = True; g1.legend.position = XL_LEGEND_POSITION.BOTTOM; g1.legend.include_in_layout = False; g1.has_title = True; g1.chart_title.text_frame.text = 'Ventas por zona'
cd2 = CategoryChartData(); cd2.categories = ['A', 'B', 'C']; cd2.add_series('Cuota', (55, 30, 15))
g2 = gs.shapes.add_chart(XL_CHART_TYPE.PIE, PI(7), PI(1.6), PI(5.6), PI(4.6), cd2).chart
g2.has_legend = True; g2.legend.position = XL_LEGEND_POSITION.RIGHT; g2.plots[0].has_data_labels = True; g2.plots[0].data_labels.show_percentage = True; g2.plots[0].data_labels.number_format = '0%'; g2.plots[0].data_labels.number_format_is_linked = False
prs.save(os.path.join(SALIDA, 'presentacion.pptx'))
print('listo')
